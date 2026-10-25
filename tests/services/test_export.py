import os
import zipfile

import pytest
from docx import Document
from docx.oxml.ns import qn

from matnyab.services import export_text

TEXT = "سلام می\u200cخواهم\n\nخط سوم ۱۲۳ abc\nپایان"


def test_txt_bom_crlf(tmp_path):
    p = tmp_path / "a" / "b" / "o.txt"
    export_text(TEXT, p, "txt")
    raw = p.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    assert b"\r\n" in raw and b"\n" not in raw.replace(b"\r\n", b"")
    assert raw.decode("utf-8-sig") == TEXT.replace("\n", "\r\n")
    assert "\u200c" in raw.decode("utf-8-sig")


def test_txt_normalises_mixed_endings_and_trailing(tmp_path):
    p = tmp_path / "o.txt"
    export_text("a\r\nb\rc\nd\n", p, "txt")
    assert p.read_bytes().decode("utf-8-sig") == "a\r\nb\r\nc\r\nd\r\n"
    export_text("", p, "txt")
    assert p.read_bytes() == b"\xef\xbb\xbf"


def test_unknown_format(tmp_path):
    with pytest.raises(ValueError):
        export_text("x", tmp_path / "o.pdf", "pdf")  # type: ignore[arg-type]
    assert not (tmp_path / "o.pdf").exists()


def test_docx_structure(tmp_path):
    p = tmp_path / "x" / "o.docx"
    export_text(TEXT, p, "docx")
    assert zipfile.is_zipfile(p)
    with zipfile.ZipFile(p) as z:
        assert "word/document.xml" in z.namelist()
        assert z.testzip() is None
    doc = Document(str(p))
    lines = TEXT.split("\n")
    assert [para.text for para in doc.paragraphs] == lines
    assert "می\u200cخواهم" in doc.paragraphs[0].text
    for para in doc.paragraphs:
        assert para._p.pPr.find(qn("w:bidi")) is not None
    for para in doc.paragraphs:
        for run in para.runs:
            rpr = run._r.rPr
            assert rpr.find(qn("w:rtl")) is not None
            assert rpr.find(qn("w:rFonts")).get(qn("w:cs")) == "Tahoma"
            assert rpr.find(qn("w:szCs")).get(qn("w:val")) == "24"
            assert rpr.find(qn("w:lang")).get(qn("w:bidi")) == "fa-IR"
    assert doc.paragraphs[1].text == "" and len(doc.paragraphs[1].runs) == 0
    s = doc.sections[0]
    assert (s.page_width.cm, s.page_height.cm) == pytest.approx((21.0, 29.7), abs=0.01)
    for m in (s.left_margin, s.right_margin, s.top_margin, s.bottom_margin):
        assert m.cm == pytest.approx(2.5, abs=0.01)
    assert doc.core_properties.title == "MatnYab export"


def test_docx_zwnj_and_special_chars_preserved(tmp_path):
    p = tmp_path / "o.docx"
    text = "\u200cنیم\u200cفاصله\u200c\tx  y\u064e\u200f\u200e"
    export_text(text, p, "docx")
    assert Document(str(p)).paragraphs[0].text == text


def test_docx_empty_and_crlf(tmp_path):
    p = tmp_path / "o.docx"
    export_text("", p, "docx")
    assert [x.text for x in Document(str(p)).paragraphs] == [""]
    export_text("a\r\nb", p, "docx")
    assert [x.text for x in Document(str(p)).paragraphs] == ["a", "b"]


def test_docx_control_chars_do_not_crash(tmp_path):
    p = tmp_path / "o.docx"
    export_text("a\x00b\x0bc", p, "docx")
    assert Document(str(p)).paragraphs[0].text == "abc"


@pytest.mark.parametrize("fmt", ["txt", "docx"])
def test_atomic_failure_keeps_old_file_no_leftovers(tmp_path, monkeypatch, fmt):
    p = tmp_path / f"o.{fmt}"
    export_text("old", p, fmt)
    before = p.read_bytes()
    monkeypatch.setattr(os, "replace", lambda *a, **k: (_ for _ in ()).throw(OSError("simulated")))
    with pytest.raises(OSError):
        export_text("new", p, fmt)
    monkeypatch.undo()
    assert p.read_bytes() == before
    assert [x.name for x in tmp_path.iterdir()] == [p.name]


def test_atomic_failure_new_target_leaves_nothing(tmp_path, monkeypatch):
    monkeypatch.setattr(os, "replace", lambda *a, **k: (_ for _ in ()).throw(OSError("simulated")))
    with pytest.raises(OSError):
        export_text("new", tmp_path / "d" / "o.txt", "txt")
    monkeypatch.undo()
    assert list((tmp_path / "d").iterdir()) == []
