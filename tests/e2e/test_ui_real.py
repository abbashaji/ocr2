"""Real window + real Tesseract engine + real file-backed services (offscreen Qt)."""

from __future__ import annotations

import threading
import zipfile
from pathlib import Path

import pytest

from matnyab.engine import scoring

pytestmark = [pytest.mark.slow, pytest.mark.ui]

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _words(text: str) -> list[str]:
    return scoring.words(text)


def test_open_recognize_copy_save_history_close(
    real_window, real_stores, clean_path, ref_text, tmp_path, monkeypatch
):
    from PySide6.QtWidgets import QApplication

    from matnyab.ui.models import ItemState

    threads_before = {t.ident for t in threading.enumerate()}
    w = real_window()  # engine=None: the window builds the real TesseractEngine itself
    w.add_paths([clean_path])
    assert w.wait_idle(180_000), "recognition did not finish"

    item = next(iter(w.items.values()))
    assert item.state == ItemState.DONE, item.state
    assert item.result is not None and item.result.mode_used in ("raw", "lex")

    # the real engine filled the editor
    editor_text = w.result.editor.toPlainText()
    assert editor_text.strip() == item.result.text.strip()
    acc = scoring.word_accuracy(ref_text, scoring.strip_markers(editor_text))
    assert acc >= 0.98, acc

    # copy
    QApplication.clipboard().clear()
    w.copy_all()
    clip = QApplication.clipboard().text()
    assert clip.strip() and _words(clip) == _words(w.final_text(editor_text))

    # save TXT: re-read with the encoding an editor would detect (BOM)
    txt = tmp_path / "out dir" / "نتیجه.txt"
    txt.parent.mkdir()
    monkeypatch.setattr(w, "_ask_save_path", lambda name: (str(txt), "txt"))
    w.save_current()
    raw = txt.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    back = raw.decode("utf-8-sig")
    assert "\r\n" in back and "\n" not in back.replace("\r\n", "")
    assert _words(back) == _words(w.final_text(editor_text))

    # save DOCX: python-docx re-reads it, paragraphs are RTL (w:bidi)
    docx_path = tmp_path / "out dir" / "نتیجه.docx"
    monkeypatch.setattr(w, "_ask_save_path", lambda name: (str(docx_path), "docx"))
    w.save_current()
    from docx import Document

    doc = Document(str(docx_path))
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    assert _words(doc_text) == _words(w.final_text(editor_text))
    with zipfile.ZipFile(docx_path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    assert "<w:bidi" in xml
    bidi_paras = [p for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.find(W_NS + "bidi") is not None]
    assert len(bidi_paras) == len(doc.paragraphs) > 0

    # history: exactly one entry for this image
    hist = real_stores[1].list()
    assert len(hist) == 1
    assert hist[0].source == clean_path.name
    assert hist[0].text == item.result.text

    # close: no leaked threads, engine released
    w.close()
    QApplication.processEvents()
    assert w.runner.holder.created is False
    leaked = [t for t in threading.enumerate()
              if t.ident not in threads_before and t.is_alive() and not t.daemon]
    leaked += [t for t in threading.enumerate()
               if t.ident not in threads_before and t.is_alive()
               and t.name.startswith("matnyab-ocr")]
    assert not leaked, [t.name for t in leaked]

