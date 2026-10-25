"""TXT and DOCX export with correct Persian right-to-left output."""

from __future__ import annotations

import io
import os
import re
from datetime import datetime, timezone

from matnyab.api import ExportFormat
from matnyab.services._atomic import atomic_write_bytes

_LINE_SPLIT = re.compile(r"\r\n|\r|\n")
# Characters that cannot appear in XML 1.0 (python-docx rejects them).
_XML_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")

DOCX_FONT = "Tahoma"
DOCX_SIZE_HALF_POINTS = 24  # 12 pt
DOCX_LANG_BIDI = "fa-IR"
DOCX_TITLE = "MatnYab export"


def _lines(text: str) -> list[str]:
    return _LINE_SPLIT.split(text)


def _txt_bytes(text: str) -> bytes:
    return "\r\n".join(_lines(text)).encode("utf-8-sig", "surrogatepass")


def _docx_bytes(text: str) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.5))

    props = doc.core_properties
    props.title = DOCX_TITLE
    props.author = "MatnYab"
    props.last_modified_by = "MatnYab"
    now = datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)
    props.created = now
    props.modified = now

    for line in _lines(_XML_ILLEGAL.sub("", text)):
        para = doc.add_paragraph()
        p_pr = para._p.get_or_add_pPr()
        p_pr.append(OxmlElement("w:bidi"))
        if not line:
            continue
        run = para.add_run(line)
        r_pr = run._r.get_or_add_rPr()
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:cs"), DOCX_FONT)
        r_pr.append(fonts)
        sz_cs = OxmlElement("w:szCs")
        sz_cs.set(qn("w:val"), str(DOCX_SIZE_HALF_POINTS))
        r_pr.append(sz_cs)
        rtl = OxmlElement("w:rtl")
        r_pr.append(rtl)
        lang = OxmlElement("w:lang")
        lang.set(qn("w:bidi"), DOCX_LANG_BIDI)
        r_pr.append(lang)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_text(text: str, path: str | os.PathLike[str], fmt: ExportFormat) -> None:
    if fmt == "txt":
        data = _txt_bytes(text)
    elif fmt == "docx":
        data = _docx_bytes(text)
    else:
        raise ValueError(f"unknown export format: {fmt!r}")
    atomic_write_bytes(path, data)
